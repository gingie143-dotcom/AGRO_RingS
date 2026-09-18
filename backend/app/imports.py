import csv
import io
import zipfile
from openpyxl import load_workbook
from pydantic import ValidationError
from .schemas import ContactIn

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000


def parse_contacts(data: bytes, filename: str):
    if len(data) > MAX_BYTES:
        raise ValueError("Максимальный размер файла — 5 МБ")
    if filename.lower().endswith(".csv"):
        text = data.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
    elif filename.lower().endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("Слишком большой распакованный XLSX")
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        rows = workbook.active.iter_rows(values_only=True)
        headers = [str(c or "").strip() for c in next(rows, [])]
        reader = (dict(zip(headers, row)) for row in rows)
    else:
        raise ValueError("Поддерживаются CSV и XLSX")
    valid, errors, seen = [], [], set()
    try:
        for index, row in enumerate(reader, 2):
            if index > MAX_ROWS + 1:
                raise ValueError("Максимум 10000 строк")
            try:
                fields = {
                    k: str(row.get(k) or "").strip()
                    for k in ["name", "phone", "position", "consent_basis"]
                }
                if any(v.startswith(("=", "@")) for v in fields.values()):
                    raise ValueError("Формулы в ячейках запрещены")
                contact = ContactIn(**fields)
                if contact.phone in seen:
                    raise ValueError("Повтор номера в файле")
                seen.add(contact.phone)
                valid.append(contact.model_dump())
            except (ValidationError, ValueError):
                errors.append(
                    {"row": index, "error": "Проверьте имя, телефон, формулы и дубликаты"}
                )
    finally:
        if filename.lower().endswith(".xlsx"):
            workbook.close()
    return {"contacts": valid, "errors": errors, "valid_count": len(valid)}
